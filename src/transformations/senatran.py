"""Build Silver fleet aggregates from SENATRAN Bronze resources."""

from __future__ import annotations

import argparse
import csv
import gzip
import io
import json
from pathlib import Path
from typing import Any, Iterable
from zipfile import ZipFile

from openpyxl import load_workbook

from src.ingestion.framework import sha256_file


BRAND_MODEL_COLUMNS = [
    "reference_period",
    "source_name",
    "source_url",
    "source_file",
    "source_hash",
    "uf",
    "municipio",
    "brand_original",
    "model_original",
    "year_manufacture",
    "quantity",
]
FUEL_COLUMNS = [
    "reference_period",
    "source_name",
    "source_url",
    "source_file",
    "source_hash",
    "uf",
    "municipio",
    "fuel_original",
    "quantity",
]
POWER_COLUMNS = [
    "reference_period",
    "source_name",
    "source_url",
    "source_file",
    "source_hash",
    "uf",
    "municipio",
    "power",
    "quantity",
]


def _metadata(path: Path) -> dict[str, Any]:
    metadata_path = path.with_name(path.name + ".metadata.json")
    if metadata_path.is_file():
        return json.loads(metadata_path.read_text(encoding="utf-8"))
    return {
        "source_name": "SENATRAN_RENAVAM",
        "source_url": "",
        "source_file": path.name,
        "file_hash": sha256_file(path),
    }


def _provenance(path: Path, period: str) -> dict[str, str]:
    metadata = _metadata(path)
    return {
        "reference_period": period,
        "source_name": str(metadata.get("source_name", "SENATRAN_RENAVAM")),
        "source_url": str(metadata.get("source_url", "")),
        "source_file": str(metadata.get("source_file", path.name)),
        "source_hash": str(metadata.get("file_hash", sha256_file(path))),
    }


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip().replace("\ufffd", "")


def _quantity(value: Any) -> int:
    return int(float(_text(value).replace(",", ".")))


def _year(value: Any) -> str:
    text = _text(value)
    if not text or text.lower() in {"sem informação", "sem informacao", "não informado", "nao informado"}:
        return ""
    try:
        return str(int(float(text)))
    except ValueError:
        return ""


def _split_brand_model(value: Any) -> tuple[str, str]:
    raw = _text(value)
    brand, separator, model = raw.partition("/")
    return brand.strip(), (model.strip() if separator else brand.strip())


def _write_rows(path: Path, columns: list[str], rows: Iterable[dict[str, Any]]) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    count = 0
    opener = gzip.open if path.suffix == ".gz" else open
    with opener(path, "wt", encoding="utf-8", newline="") as file_obj:
        writer = csv.DictWriter(file_obj, fieldnames=columns)
        writer.writeheader()
        for row in rows:
            writer.writerow({column: row.get(column, "") for column in columns})
            count += 1
    return count


def _brand_model_rows(source_path: Path, period: str) -> Iterable[dict[str, Any]]:
    provenance = _provenance(source_path, period)
    with ZipFile(source_path) as archive:
        member = archive.namelist()[0]
        with archive.open(member) as binary_stream:
            text_stream = io.TextIOWrapper(binary_stream, encoding="utf-8-sig", newline="")
            reader = csv.reader(text_stream, delimiter=";")
            next(reader, None)
            for row in reader:
                if len(row) < 5 or not _text(row[0]):
                    continue
                brand, model = _split_brand_model(row[2])
                yield {
                    **provenance,
                    "uf": _text(row[0]),
                    "municipio": _text(row[1]),
                    "brand_original": brand,
                    "model_original": model,
                    "year_manufacture": _year(row[3]),
                    "quantity": _quantity(row[4]),
                }


def _xlsx_rows(source_path: Path, period: str, *, measure: str) -> Iterable[dict[str, Any]]:
    provenance = _provenance(source_path, period)
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            next(rows, None)
            for row in rows:
                if len(row) < 4 or not _text(row[0]):
                    continue
                record = {
                    **provenance,
                    "uf": _text(row[0]),
                    "municipio": _text(row[1]),
                    "quantity": _quantity(row[3]),
                }
                if measure == "fuel":
                    record["fuel_original"] = _text(row[2])
                else:
                    record["power"] = _text(row[2])
                yield record
    finally:
        workbook.close()


def build_silver_from_bronze(
    period: str,
    *,
    bronze_root: Path = Path("data/bronze/senatran"),
    silver_root: Path = Path("data/silver/senatran"),
) -> dict[str, Any]:
    """Stream all three official SENATRAN aggregates into canonical Silver CSVs."""

    bronze_period = Path(bronze_root) / period
    auxiliary = bronze_period / "auxiliary"
    brand_model_files = list(bronze_period.glob("*.zip"))
    fuel_files = list(auxiliary.glob("*COMBUSTIVEL*.xlsx"))
    power_files = list(auxiliary.glob("*POTENCIA*.xlsx"))
    if not brand_model_files or not fuel_files or not power_files:
        raise FileNotFoundError("Expected SENATRAN brand/model ZIP and fuel/power XLSX files were not found")

    output_dir = Path(silver_root) / period
    brand_model_path = output_dir / "fleet_brand_model.csv.gz"
    fuel_path = output_dir / "fleet_fuel.csv"
    power_path = output_dir / "fleet_power.csv"
    brand_model_count = _write_rows(brand_model_path, BRAND_MODEL_COLUMNS, _brand_model_rows(brand_model_files[0], period))
    fuel_count = _write_rows(fuel_path, FUEL_COLUMNS, _xlsx_rows(fuel_files[0], period, measure="fuel"))
    power_count = _write_rows(power_path, POWER_COLUMNS, _xlsx_rows(power_files[0], period, measure="power"))

    report = {
        "status": "PASS",
        "reference_period": period,
        "outputs": {
            "fleet_brand_model": {"path": str(brand_model_path), "rows": brand_model_count},
            "fleet_fuel": {"path": str(fuel_path), "rows": fuel_count},
            "fleet_power": {"path": str(power_path), "rows": power_count},
        },
        "grain": {
            "fleet_brand_model": "UF + município + marca/modelo + ano de fabricação",
            "fleet_fuel": "UF + município + combustível",
            "fleet_power": "UF + município + potência",
        },
        "limitations": [
            "SENATRAN publishes these aggregates separately; fuel and power are not joined to an individual vehicle or brand/model row.",
            "SENATRAN represents fleet stock, not monthly registrations or sales flow.",
        ],
    }
    report_path = output_dir / "quality_report.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description="Build Silver SENATRAN fleet aggregates.")
    parser.add_argument("--period", required=True)
    parser.add_argument("--bronze-root", type=Path, default=Path("data/bronze/senatran"))
    parser.add_argument("--silver-root", type=Path, default=Path("data/silver/senatran"))
    args = parser.parse_args()
    report = build_silver_from_bronze(args.period, bronze_root=args.bronze_root, silver_root=args.silver_root)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
